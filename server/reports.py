"""Excel sales report.

Built on the server rather than in the browser so it reads every bill in the
range straight from the database. The analytics page only holds the first
thousand bills it fetched; a report built from that would quietly undercount a
busy month.

Look: Times New Roman, black only, a thin border on every table cell and a
double rule above each total - a plain, printable ledger.
"""

import math
from collections import defaultdict
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from server.timeutil import LOCAL_TZ, as_utc

RESTAURANT = "S&H Arabian Mandi Restaurant"

# --- styling ---------------------------------------------------------------

FONT = "Times New Roman"
BLACK = "000000"

# The trailing "_)" leaves a small space after each number, so right-aligned
# figures never touch the cell's line (works in Excel and LibreOffice alike).
MONEY = '"₹"#,##0.00_)'
INTEGER = "#,##0_)"
PERCENT = "0.0%_)"
DATE = "dd-mmm-yyyy"
TIME = "hh:mm AM/PM"

TITLE_FONT = Font(name=FONT, size=16, bold=True, color=BLACK)
SUBTITLE_FONT = Font(name=FONT, size=11, italic=True, color=BLACK)
HEADER_FONT = Font(name=FONT, size=11, bold=True, color=BLACK)
BODY_FONT = Font(name=FONT, size=11, color=BLACK)
BOLD_FONT = Font(name=FONT, size=11, bold=True, color=BLACK)

THIN = Side(style="thin", color=BLACK)
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=Side(style="medium", color=BLACK))
TOTAL_BORDER = Border(left=THIN, right=THIN, top=Side(style="double", color=BLACK), bottom=THIN)

LINE_HEIGHT = 15      # points per line of 11pt text
ROW_HEIGHT = 18       # one-line table row, with a little air above and below
FILTER_ARROW = 3      # characters the filter button covers in a header cell
# Sheets narrow enough to print upright; the wide ones print landscape.
LANDSCAPE = {"Daily Sales", "Bill Register"}

# Payment methods the report gives their own column; anything else is "Other".
PAYMENT_COLUMNS = ("Cash", "UPI", "Card")
# The older till records every bill as "Cash / UPI"; count it as cash.
CASH_ALIASES = {"Cash", "Cash / UPI", "Cash/UPI"}


# --- bill helpers ----------------------------------------------------------
#
# Bills in the collection come in more than one shape: the current backend
# writes billNumber/subtotal/discountAmount/createdBy, older builds wrote only
# token/total. Every read goes through these so the report handles both.


def _local(bill: dict) -> datetime:
    return as_utc(bill["createdAt"]).astimezone(LOCAL_TZ)


def _number(bill: dict):
    return bill.get("billNumber") or bill.get("token") or ""


def _total(bill: dict) -> float:
    return float(bill.get("total") or 0)


def _subtotal(bill: dict) -> float:
    value = bill.get("subtotal")
    return float(value) if value is not None else _total(bill)


def _discount(bill: dict) -> float:
    value = bill.get("discountAmount")
    if value is not None:
        return float(value)
    return max(0.0, _subtotal(bill) - _total(bill))


def _payment_bucket(payment: str) -> str:
    if payment in CASH_ALIASES:
        return "Cash"
    return payment if payment in PAYMENT_COLUMNS else "Other"


# --- sheet helpers ---------------------------------------------------------


def _shown_width(value, fmt: str | None) -> int:
    """Roughly how many characters a value takes once Excel formats it."""
    if value is None or value == "":
        return 0
    if isinstance(value, str):  # labels such as "TOTAL" in a number column
        return len(value)
    if fmt == MONEY:
        return len(f"₹{float(value):,.2f}") + 1  # the ₹ sign prints wide
    if fmt == PERCENT:
        return len(f"{float(value) * 100:.1f}%")
    if fmt == INTEGER:
        return len(f"{int(value):,}")
    if fmt == DATE:
        return 11
    if fmt == TIME:
        return 8
    return len(str(value))


def _align(value, fmt: str | None = None, wrap: bool = False) -> Alignment:
    """Amounts to the right, counts, dates and times centred, text to the
    left with a small margin, so nothing touches the cell's lines."""
    if fmt == INTEGER and not isinstance(value, str):
        return Alignment(horizontal="center", vertical="top")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return Alignment(horizontal="right", vertical="top")
    if isinstance(value, (date, time)):
        return Alignment(horizontal="center", vertical="top")
    return Alignment(horizontal="left", vertical="top", indent=1, wrap_text=wrap)


def _title(ws, title: str, subtitle: str) -> int:
    """Two-line heading; returns the table's header row.

    Not merged: the text runs on over the empty cells to its right, so a long
    title is never clipped to the width of the first columns.
    """
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[1].height = 24
    return 4


def _header(ws, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=column, value=label)
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = LINE_HEIGHT * 2


def _body(ws, first_row: int, rows: list[list], formats: list[str | None]) -> int:
    """Data rows with per-column number formats; returns the next free row."""
    row = first_row
    for values in rows:
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.font = BODY_FONT
            cell.border = GRID
            fmt = formats[column - 1]
            cell.alignment = _align(value, fmt)
            if fmt:
                cell.number_format = fmt
        ws.row_dimensions[row].height = ROW_HEIGHT
        row += 1
    return row


def _totals(ws, row: int, values: list, formats: list[str | None]) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = BOLD_FONT
        cell.border = TOTAL_BORDER
        fmt = formats[column - 1]
        cell.alignment = _align(value, fmt)
        if fmt and value not in ("", None):
            cell.number_format = fmt
    ws.row_dimensions[row].height = ROW_HEIGHT + 2


def _fit_columns(ws, header_row: int, labels: list[str], rows: list[list],
                 formats: list[str | None], fixed: dict[int, int] | None = None) -> None:
    """Size every column to its widest heading or value (plus room for the
    filter arrow), so nothing shows as #### or hides behind the arrow."""
    fixed = fixed or {}
    for index, label in enumerate(labels):
        column = index + 1
        if column in fixed:
            width = fixed[column]
        else:
            longest_word = max(len(word) for word in label.split())
            width = max(longest_word + FILTER_ARROW, len(label) // 2 + FILTER_ARROW)
            for values in rows:
                width = max(width, _shown_width(values[index], formats[index]) + 4)
            width = min(width, 45)
        ws.column_dimensions[get_column_letter(column)].width = width + 1


def _finish_table(ws, header_row: int, last_row: int, width: int) -> None:
    """Freeze the header and add filter arrows over the data."""
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"
    if last_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(width)}{last_row}"


# --- the sheets ------------------------------------------------------------


def _summary(ws, bills: list[dict], period: str, generated: str) -> None:
    ws.title = "Summary"
    header_row = _title(ws, f"{RESTAURANT} — Sales Report", f"{period}     Generated {generated}")

    revenue = sum(_total(b) for b in bills)
    gross = sum(_subtotal(b) for b in bills)
    discounts = sum(_discount(b) for b in bills)
    items_sold = sum(int(i.get("qty") or 0) for b in bills for i in b.get("items") or [])
    days = len({_local(b).date() for b in bills})

    kpis = [
        ["Bills", len(bills)],
        ["Revenue", revenue],
        ["Gross (before discount)", gross],
        ["Discounts given", discounts],
        ["Average bill", revenue / len(bills) if bills else 0],
        ["Items sold", items_sold],
        ["Trading days", days],
        ["Average per day", revenue / days if days else 0],
    ]
    kpi_formats = [INTEGER, MONEY, MONEY, MONEY, MONEY, INTEGER, INTEGER, MONEY]

    _header(ws, header_row, ["Particulars", "Value"])
    ws.row_dimensions[header_row].height = LINE_HEIGHT + 5
    for offset, ((label, value), fmt) in enumerate(zip(kpis, kpi_formats), start=1):
        row = header_row + offset
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = BOLD_FONT
        label_cell.border = GRID
        label_cell.alignment = _align(label)
        value_cell = ws.cell(row=row, column=2, value=value)
        value_cell.font = BODY_FONT
        value_cell.border = GRID
        value_cell.number_format = fmt
        value_cell.alignment = _align(value)
        ws.row_dimensions[row].height = ROW_HEIGHT + 2

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22


def _daily(ws, bills: list[dict], period: str) -> None:
    ws.title = "Daily Sales"
    labels = ["S.No", "Date", "Day", "Bills", *PAYMENT_COLUMNS, "Other", "Discounts", "Total", "Average bill"]
    formats = [INTEGER, DATE, None, INTEGER, MONEY, MONEY, MONEY, MONEY, MONEY, MONEY, MONEY]
    header_row = _title(ws, "Daily Sales", period)
    _header(ws, header_row, labels)

    by_day: dict[date, list[dict]] = defaultdict(list)
    for bill in bills:
        by_day[_local(bill).date()].append(bill)

    rows = []
    for index, day in enumerate(sorted(by_day), start=1):
        day_bills = by_day[day]
        buckets = defaultdict(float)
        for bill in day_bills:
            buckets[_payment_bucket(bill.get("payment") or "")] += _total(bill)
        total = sum(_total(b) for b in day_bills)
        rows.append([
            index,
            day,
            day.strftime("%a"),
            len(day_bills),
            *(buckets[p] for p in PAYMENT_COLUMNS),
            buckets["Other"],
            sum(_discount(b) for b in day_bills),
            total,
            total / len(day_bills),
        ])

    last = _body(ws, header_row + 1, rows, formats) - 1

    totals = []
    if rows:
        count = sum(r[3] for r in rows)
        grand = sum(r[9] for r in rows)
        totals = ["TOTAL", "", "", count,
                  *(sum(r[4 + i] for r in rows) for i in range(len(PAYMENT_COLUMNS))),
                  sum(r[7] for r in rows), sum(r[8] for r in rows), grand, grand / count if count else 0]
        _totals(ws, last + 1, totals, formats)

    _fit_columns(ws, header_row, labels, rows + ([totals] if totals else []), formats)
    _finish_table(ws, header_row, last, len(labels))


def _payments(ws, bills: list[dict], period: str) -> None:
    ws.title = "Payment Methods"
    labels = ["Payment method", "Bills", "Amount", "Share"]
    formats = [None, INTEGER, MONEY, PERCENT]
    header_row = _title(ws, "Payment Methods", period)
    _header(ws, header_row, labels)

    totals: dict[str, list] = defaultdict(lambda: [0, 0.0])
    for bill in bills:
        entry = totals[bill.get("payment") or "Unknown"]
        entry[0] += 1
        entry[1] += _total(bill)

    revenue = sum(amount for _, amount in totals.values()) or 1
    rows = [
        [method, count, amount, amount / revenue]
        for method, (count, amount) in sorted(totals.items(), key=lambda kv: -kv[1][1])
    ]
    last = _body(ws, header_row + 1, rows, formats) - 1
    total_row = []
    if rows:
        total_row = ["TOTAL", sum(r[1] for r in rows), sum(r[2] for r in rows), 1]
        _totals(ws, last + 1, total_row, formats)

    _fit_columns(ws, header_row, labels, rows + ([total_row] if total_row else []), formats)
    _finish_table(ws, header_row, last, len(labels))


def _order_types(ws, bills: list[dict], period: str) -> None:
    ws.title = "Order Types"
    labels = ["Order type", "Bills", "Amount", "Average bill"]
    formats = [None, INTEGER, MONEY, MONEY]
    header_row = _title(ws, "Order Types", period)
    _header(ws, header_row, labels)

    totals: dict[str, list] = defaultdict(lambda: [0, 0.0])
    for bill in bills:
        entry = totals[bill.get("orderType") or "Unknown"]
        entry[0] += 1
        entry[1] += _total(bill)

    rows = [
        [kind, count, amount, amount / count if count else 0]
        for kind, (count, amount) in sorted(totals.items(), key=lambda kv: -kv[1][1])
    ]
    last = _body(ws, header_row + 1, rows, formats) - 1
    total_row = []
    if rows:
        count = sum(r[1] for r in rows)
        amount = sum(r[2] for r in rows)
        total_row = ["TOTAL", count, amount, amount / count if count else 0]
        _totals(ws, last + 1, total_row, formats)

    _fit_columns(ws, header_row, labels, rows + ([total_row] if total_row else []), formats)
    _finish_table(ws, header_row, last, len(labels))


def _items(ws, bills: list[dict], period: str) -> None:
    ws.title = "Item Sales"
    labels = ["Rank", "Item", "Qty sold", "Revenue", "Share of revenue"]
    formats = [INTEGER, None, INTEGER, MONEY, PERCENT]
    header_row = _title(ws, "Item Sales", period)
    _header(ws, header_row, labels)

    totals: dict[str, list] = defaultdict(lambda: [0, 0.0])
    for bill in bills:
        for item in bill.get("items") or []:
            qty = int(item.get("qty") or 0)
            entry = totals[item.get("name") or "Unnamed"]
            entry[0] += qty
            entry[1] += float(item.get("price") or 0) * qty

    revenue = sum(amount for _, amount in totals.values()) or 1
    ordered = sorted(totals.items(), key=lambda kv: (-kv[1][1], kv[0]))
    rows = [
        [rank, name, qty, amount, amount / revenue]
        for rank, (name, (qty, amount)) in enumerate(ordered, start=1)
    ]
    last = _body(ws, header_row + 1, rows, formats) - 1
    total_row = []
    if rows:
        total_row = ["", "TOTAL", sum(r[2] for r in rows), sum(r[3] for r in rows), 1]
        _totals(ws, last + 1, total_row, formats)

    _fit_columns(ws, header_row, labels, rows + ([total_row] if total_row else []), formats)
    _finish_table(ws, header_row, last, len(labels))


def _register(ws, bills: list[dict], period: str) -> None:
    ws.title = "Bill Register"
    labels = ["Bill No", "Date", "Time", "Order type", "Table", "Items", "Item details",
              "Subtotal", "Discount", "Total", "Payment", "Billed by"]
    formats = [None, DATE, TIME, None, None, INTEGER, None,
               MONEY, MONEY, MONEY, None, None]
    details_col = 7
    details_width = 50
    header_row = _title(ws, "Bill Register", period)
    _header(ws, header_row, labels)

    rows = []
    for bill in sorted(bills, key=_local):
        # Excel has no timezones, so write naive IST; seconds are noise on a bill.
        when = _local(bill).replace(tzinfo=None, second=0, microsecond=0)
        items = bill.get("items") or []
        rows.append([
            _number(bill),
            when.date(),
            time(when.hour, when.minute),
            bill.get("orderType") or "",
            bill.get("tableNo") or "",
            sum(int(i.get("qty") or 0) for i in items),
            ", ".join(f"{i.get('qty', 0)} × {i.get('name', '')}" for i in items),
            _subtotal(bill),
            _discount(bill),
            _total(bill),
            bill.get("payment") or "",
            bill.get("createdBy") or "",
        ])

    last = _body(ws, header_row + 1, rows, formats) - 1
    # Long item lists wrap inside their cell; give each row the height its
    # text needs (Excel and LibreOffice do not grow rows for wrapped text
    # written by a program).
    chars_per_line = int((details_width - 2) * 1.25)  # Times New Roman runs narrow
    for offset, values in enumerate(rows):
        row = header_row + 1 + offset
        ws.cell(row=row, column=details_col).alignment = _align(values[details_col - 1], wrap=True)
        lines = max(1, math.ceil(len(values[details_col - 1]) / chars_per_line))
        ws.row_dimensions[row].height = max(ROW_HEIGHT, LINE_HEIGHT * lines + 3)

    total_row = []
    if rows:
        total_row = ["TOTAL", "", "", "", "", sum(r[5] for r in rows), "",
                     sum(r[7] for r in rows), sum(r[8] for r in rows), sum(r[9] for r in rows), "", ""]
        _totals(ws, last + 1, total_row, formats)

    _fit_columns(ws, header_row, labels, rows + ([total_row] if total_row else []), formats,
                 fixed={details_col: details_width})
    _finish_table(ws, header_row, last, len(labels))


# --- entry point -----------------------------------------------------------


def build_sales_workbook(bills: list[dict], date_from: date, date_to: date) -> bytes:
    """Render the report for ``bills`` (already filtered to the range) as .xlsx bytes."""
    period = f"{date_from:%d %b %Y} to {date_to:%d %b %Y}"
    generated = datetime.now(LOCAL_TZ).strftime("%d %b %Y, %I:%M %p")

    wb = Workbook()
    _summary(wb.active, bills, period, generated)
    _daily(wb.create_sheet(), bills, period)
    _payments(wb.create_sheet(), bills, period)
    _order_types(wb.create_sheet(), bills, period)
    _items(wb.create_sheet(), bills, period)
    _register(wb.create_sheet(), bills, period)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False  # the tables carry their own lines
        sheet.page_setup.orientation = "landscape" if sheet.title in LANDSCAPE else "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins.left = sheet.page_margins.right = 0.4
        sheet.oddFooter.center.text = "Page &P of &N"
        sheet.oddFooter.center.font = f"{FONT},Regular"
        sheet.oddFooter.center.size = 9

    wb.properties.title = f"{RESTAURANT} sales {period}"
    wb.properties.creator = RESTAURANT

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
